"""Turn an uploaded file (.txt, .md, .pdf, .docx) into plain text."""
import io
import re

MAX_CHARS = 40_000  # keeps one prompt a sensible size; longer files are cut with a warning


def extract_text(name: str, data: bytes) -> tuple[str, bool]:
    """Return (text, was_truncated). Raises ValueError with a readable message."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""

    if ext in ("txt", "md"):
        text = data.decode("utf-8", errors="replace")

    elif ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)

    elif ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many process docs keep steps in tables
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)

    else:
        raise ValueError("Unsupported file type. Use .txt, .md, .pdf or .docx.")

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 30:
        raise ValueError(
            "No readable text found. If this is a scanned PDF (a photo of a page), "
            "it needs OCR (text recognition) first."
        )
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS], True
    return text, False
